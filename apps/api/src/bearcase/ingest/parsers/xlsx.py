"""XLSX parsing with openpyxl in read-only, data-only mode. Formulas are read as cached values
and never evaluated; VBA, external links, and embedded objects are ignored."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from bearcase.ingest.parsers.common import ParsedChunk, ParseError, ParseResult


def _cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def parse_xlsx(data: bytes, max_rows: int = 20000) -> ParseResult:
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise ParseError(f"Could not open workbook: {exc}") from exc
    chunks: list[ParsedChunk] = []
    texts: list[str] = []
    total_rows = 0
    sheets: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        sheet_rows = 0
        header: list[str] = []
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            total_rows += 1
            if total_rows > max_rows:
                raise ParseError(f"Workbook exceeds the {max_rows} row limit.")
            values = [_cell_text(v) for v in row]
            if not any(values):
                continue
            sheet_rows += 1
            if not header:
                header = values
            cells: dict[str, Any] = {}
            for col_idx, raw in enumerate(row, start=1):
                if raw is None or raw == "":
                    continue
                cells[f"{get_column_letter(col_idx)}{row_idx}"] = raw if isinstance(raw, int | float) else str(raw)
            row_text = " | ".join(v for v in values if v)
            chunks.append(
                ParsedChunk(
                    kind="sheet_row",
                    locator={"sheet": ws.title, "row": row_idx, "range": f"A{row_idx}:{get_column_letter(len(values))}{row_idx}"},
                    text=f"{ws.title}!{row_idx}: {row_text}",
                    structured={"cells": cells, "values": values, "header": header},
                )
            )
            texts.append(row_text)
        sheets.append({"title": ws.title, "rows": sheet_rows})
    wb.close()
    return ParseResult(
        chunks=chunks, sheet_count=len(sheets), row_count=total_rows, full_text="\n".join(texts), meta={"sheets": sheets}
    )
