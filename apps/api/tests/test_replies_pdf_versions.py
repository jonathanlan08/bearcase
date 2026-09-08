"""Seller replies, the one-page PDF, and revised-document versions with a diff."""

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from bearcase.api.app import app


def _own_demo():
    c = TestClient(app)
    return c, c.post("/api/demo/session").json()["id"]


def test_seller_replies_round_trip_into_the_export():
    c, deal = _own_demo()
    q = c.get(f"/api/deals/{deal}/seller-questions").json()["questions"][1]
    r = c.post(f"/api/deals/{deal}/seller-replies", json={"question_id": q["id"], "question_text": q["question"], "reply_text": "We use bookings, not revenue, for the growth figure.", "outcome": "dodged"})
    assert r.status_code == 201, r.text
    assert r.json()["outcome"] == "dodged"
    listed = c.get(f"/api/deals/{deal}/seller-replies").json()
    assert listed["totals"]["dodged"] == 1 and q["id"] in listed["replies"]
    r2 = c.post(f"/api/deals/{deal}/seller-replies", json={"question_id": q["id"], "question_text": q["question"], "reply_text": "Sent the monthly revenue ledger.", "outcome": "answered"})
    assert r2.status_code == 201
    listed = c.get(f"/api/deals/{deal}/seller-replies").json()
    assert listed["totals"] == {"answered": 1, "dodged": 0, "needs_document": 0}, "the latest reply per question wins"
    md = c.get(f"/api/deals/{deal}/seller-questions/export?format=md").text
    assert "Seller's reply (Answered): Sent the monthly revenue ledger." in md
    assert c.post(f"/api/deals/{deal}/seller-replies", json={"question_id": "nope", "question_text": "x", "reply_text": "y", "outcome": "answered"}).status_code == 422


def test_one_page_summary_pdf():
    c, deal = _own_demo()
    r = c.get(f"/api/deals/{deal}/summary.pdf")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/pdf")
    assert r.content[:4] == b"%PDF" and len(r.content) > 2000
    assert "summary.pdf" in r.headers["content-disposition"]
    assert any(e["event_type"] == "summary.exported" for e in c.get(f"/api/deals/{deal}/audit").json())


def _revised_statement(c: TestClient, deal: str) -> tuple[str, bytes]:
    docs = c.get(f"/api/deals/{deal}/documents").json()
    fs = next(d for d in docs if d["doc_type"] == "financial_statements")
    # rebuild the workbook with FY2024 revenue restated
    from bearcase.fixtures.generate import build_financial_statements

    wb = load_workbook(BytesIO(build_financial_statements()))
    ws = wb["Income Statement"]
    for row in ws.iter_rows(min_row=4):
        if row[0].value == "Revenue":
            row[3].value = 13_100_000
    buf = BytesIO()
    wb.save(buf)
    return fs["id"], buf.getvalue()


def test_new_version_and_diff():
    c, deal = _own_demo()
    doc_id, revised = _revised_statement(c, deal)
    assert c.get(f"/api/deals/{deal}/documents/{doc_id}/diff").json()["comparable"] is False
    r = c.post(f"/api/deals/{deal}/documents/{doc_id}/versions", files={"file": ("northstar-financial-statements-restated.xlsx", revised, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 201, r.text
    assert r.json()["version"]["version_no"] == 2
    d = c.get(f"/api/deals/{deal}/documents/{doc_id}/diff").json()
    assert d["comparable"] and d["kind"] == "statement" and d["old_version"] == 1 and d["new_version"] == 2
    assert d["changes"] == [{"period": "FY2024", "line_key": "revenue", "before": "12950000", "after": "13100000", "cell": "D4", "dependents": d["changes"][0]["dependents"]}]
    assert "cagr" in d["changes"][0]["dependents"] and d["stale"] is True
    assert any(cl["metric_key"] == "revenue" for cl in d["claims_affected"])
    # same file again is refused; a different kind of file is refused
    assert c.post(f"/api/deals/{deal}/documents/{doc_id}/versions", files={"file": ("again.xlsx", revised, "application/octet-stream")}).status_code == 422
    assert c.post(f"/api/deals/{deal}/documents/{doc_id}/versions", files={"file": ("x.csv", b"customer_name,revenue_type,revenue\\nA,maintenance,1\\n", "text/csv")}).status_code == 422
