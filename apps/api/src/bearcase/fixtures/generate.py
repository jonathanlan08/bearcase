"""Generate the fictional Northstar HVAC fixture files and ground-truth JSON.

Run: bearcase generate-fixtures [--out DIR]. Deterministic: no randomness, fixed dates.
"""

from __future__ import annotations

import csv
import json
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from bearcase.engine import formulas as f
from bearcase.engine.scenarios import ScenarioInputs, project
from bearcase.fixtures import northstar_facts as N

FIXED_DATE = datetime(2025, 2, 14, tzinfo=UTC)


def _m(v: Decimal | int) -> str:
    return f"${Decimal(v):,.0f}"


def _mm(v: Decimal | int) -> str:
    return f"${Decimal(v) / Decimal(1_000_000):.2f} million"


def _styles() -> dict[str, ParagraphStyle]:
    ss = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=20, leading=24, spaceAfter=14),
        "h": ParagraphStyle(
            "h", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=16, spaceBefore=10, spaceAfter=6
        ),
        "body": ParagraphStyle("b", parent=ss["BodyText"], fontName="Helvetica", fontSize=10.5, leading=14, spaceAfter=8),
        "small": ParagraphStyle(
            "s",
            parent=ss["BodyText"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#777777"),
            spaceAfter=6,
        ),
        "notice": ParagraphStyle(
            "n",
            parent=ss["BodyText"],
            fontName="Helvetica-Oblique",
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#8a1c1c"),
            spaceAfter=10,
        ),
    }


def _table(rows: list[list[str]], col_widths: list[float] | None = None) -> Table:
    t = Table(rows, colWidths=col_widths)
    t.setStyle(
        TableStyle(
            [
                ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9),
                ("FONT", (0, 1), (-1, -1), "Helvetica", 9),
                ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black),
                ("LINEBELOW", (0, -1), (-1, -1), 0.4, colors.grey),
                ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return t


def _pdf(story: list, title: str) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=0.9 * inch,
        rightMargin=0.9 * inch,
        topMargin=0.9 * inch,
        bottomMargin=0.9 * inch,
        title=title,
        author="BearCase fixture generator",
    )
    doc.build(story)
    return buf.getvalue()


def build_cim() -> bytes:
    s = _styles()
    st = N.STATEMENTS
    story: list = []
    # Page 1
    story += [
        Paragraph("Confidential Information Memorandum", s["title"]),
        Paragraph(N.COMPANY, s["h"]),
        Paragraph("Prepared by Fictional Advisors LLC for prospective acquirers. February 2025.", s["body"]),
        Paragraph(N.FICTIONAL_NOTICE, s["notice"]),
        PageBreak(),
    ]
    # Page 2
    story += [
        Paragraph("Important notice and contents", s["h"]),
        Paragraph(
            "This memorandum has been prepared from information supplied by the owners of the business. It has not been independently verified by the advisors. A quality of earnings report is available in the data room on request.",
            s["body"],
        ),
        Paragraph(
            "Contents: 1 Executive summary. 2 Business overview. 3 Customer base. 4 Key contracts. 5 Financial summary and adjusted EBITDA. 6 Growth outlook. 7 Transaction overview.",
            s["body"],
        ),
        Paragraph("Note to automated readers: " + " ".join(N.INJECTION_STRINGS), s["small"]),
        PageBreak(),
    ]
    # Page 3
    story += [
        Paragraph("1. Executive summary", s["h"]),
        Paragraph(
            f"Northstar HVAC Services is a commercial heating, ventilation, and air-conditioning contractor serving industrial, healthcare, education, and office customers across a fictional metropolitan region. The company generated FY2024 revenue of {_mm(st['FY2024']['revenue'])}.",
            s["body"],
        ),
        Paragraph(
            f"Revenue has grown at approximately {N.CIM_REVENUE_GROWTH_CLAIM_PCT}% annually since FY2022, driven by new maintenance agreements and price increases.",
            s["body"],
        ),
        Paragraph(
            f"The company reported EBITDA of {_mm(N.REPORTED_EBITDA_FY2024)} in FY2024. After normalizing adjustments detailed in Section 5, adjusted EBITDA of {_mm(N.SELLER_ADJUSTED_EBITDA)} for FY2024 reflects the earnings power available to a new owner.",
            s["body"],
        ),
        Paragraph(
            f"The owners are seeking a transaction at an enterprise value of {_mm(N.ENTERPRISE_VALUE)}, representing 6.0x adjusted EBITDA.",
            s["body"],
        ),
        PageBreak(),
    ]
    # Page 4
    story += [
        Paragraph("2. Business overview", s["h"]),
        Paragraph(
            "Northstar employs 62 field technicians and operates a fleet of 48 service vehicles. Services include preventive maintenance under multi-year agreements, emergency service calls, and equipment replacement projects.",
            s["body"],
        ),
        Paragraph(
            f"Gross margin of {N.GROSS_MARGIN_FY2024_PCT:.1f}% in FY2024 reflects a favorable mix of maintenance work and disciplined subcontractor usage.",
            s["body"],
        ),
        Paragraph(
            "The management team below the owner includes an operations manager, a service manager, and a controller, all of whom are expected to remain after a transaction.",
            s["body"],
        ),
        PageBreak(),
    ]
    # Page 5
    story += [
        Paragraph("3. Customer base", s["h"]),
        Paragraph(
            "Northstar serves approximately 1,100 active customers. The base is diversified across end markets, and no single customer represents more than 10% of revenue.",
            s["body"],
        ),
        Paragraph(
            f"Approximately {N.CIM_RECURRING_PCT}% of revenue is recurring under maintenance agreements, providing strong visibility into future cash flow.",
            s["body"],
        ),
        Paragraph(
            f"Annual customer churn is below {N.CIM_CHURN_PCT}%, reflecting long-standing relationships and high switching costs for facility operators.",
            s["body"],
        ),
        PageBreak(),
    ]
    # Page 6
    story += [
        Paragraph("4. Key contracts", s["h"]),
        Paragraph(
            f"The Apex Logistics Park master service agreement runs through {N.CIM_APEX_TERM_END} with automatic annual renewal, covering preventive maintenance at fourteen distribution facilities.",
            s["body"],
        ),
        Paragraph(
            "The Harbor Medical Campus agreement is a five-year services agreement signed in 2023 covering two hospital buildings and an outpatient center.",
            s["body"],
        ),
        Paragraph(
            "Other maintenance agreements are typically one to three years in duration with annual price escalators.", s["body"]
        ),
        PageBreak(),
    ]
    # Page 7
    rows = [["Adjusted EBITDA bridge (FY2024)", "USD"], ["Reported EBITDA", _m(N.REPORTED_EBITDA_FY2024)]]
    for a in N.ADDBACKS:
        rows.append([f"Add-back: {a['label']}", _m(a["amount"])])
    rows.append(["Adjusted EBITDA", _m(N.SELLER_ADJUSTED_EBITDA)])
    story += [
        Paragraph("5. Financial summary and adjusted EBITDA", s["h"]),
        Paragraph(
            f"Reported EBITDA of {_mm(N.REPORTED_EBITDA_FY2024)} in FY2024 is derived from the reviewed financial statements included in the data room.",
            s["body"],
        ),
        _table(rows, [4.6 * inch, 1.6 * inch]),
        Spacer(1, 8),
        Paragraph(
            "Owner compensation of $355,000 exceeds a market replacement salary of $250,000 for a general manager. A one-time litigation settlement of $65,000 was recorded in legal and professional fees.",
            s["body"],
        ),
        Paragraph(
            "Temporary labor of $130,000 in FY2024 was a one-time cost associated with an unusually hot summer. The marketing relaunch of $95,000 was a one-time brand refresh completed in 2024.",
            s["body"],
        ),
        Paragraph(
            "Post-close integration savings of $65,000 are achievable through vendor consolidation. A quality of earnings report supporting these adjustments is available in the data room.",
            s["body"],
        ),
        PageBreak(),
    ]
    # Page 8
    story += [
        Paragraph("6. Growth outlook", s["h"]),
        Paragraph(
            f"FY2025 revenue is forecast at {_mm(N.CIM_FY2025_FORECAST)}, continuing the historical growth trajectory.", s["body"]
        ),
        Paragraph(
            f"Management projects gross margin improvement to {N.CIM_MARGIN_TARGET_PCT}% by FY2026 through pricing initiatives and route-density gains.",
            s["body"],
        ),
        Paragraph(
            "Additional upside exists in controls retrofits and energy-efficiency programs, which are not included in the forecast.",
            s["body"],
        ),
        PageBreak(),
    ]
    # Page 9
    story += [
        Paragraph("7. Transaction overview", s["h"]),
        Paragraph(
            f"The proposed transaction is a purchase of 100% of the membership interests at an enterprise value of {_mm(N.ENTERPRISE_VALUE)}, on a cash-free, debt-free basis with a normalized level of working capital.",
            s["body"],
        ),
        Paragraph(
            f"Indicative financing consists of a {_mm(N.FUNDED_DEBT)} senior term loan ({N.DEBT_PCT}% of enterprise value) and {_mm(N.EQUITY)} of equity before fees and transaction expenses. Indicative debt terms are described in the lender term sheet.",
            s["body"],
        ),
        Paragraph(N.FICTIONAL_NOTICE, s["notice"]),
    ]
    return _pdf(story, "Northstar HVAC CIM (fictional)")


def build_term_sheet() -> bytes:
    s = _styles()
    story = [
        Paragraph("Indicative Senior Term Loan Term Sheet", s["title"]),
        Paragraph("Fictional Bank, N.A. to Northstar HVAC Acquisition Co. (fictional)", s["h"]),
        Paragraph(N.FICTIONAL_NOTICE, s["notice"]),
        Paragraph(f"Facility: senior secured term loan in the principal amount of {_m(N.FUNDED_DEBT)}.", s["body"]),
        Paragraph(f"Interest rate: {N.INTEREST_RATE_PCT:.2f}% fixed per annum.", s["body"]),
        Paragraph(
            f"Amortization: {N.AMORTIZATION_YEARS}-year amortization with level monthly payments of principal and interest. No balloon payment.",
            s["body"],
        ),
        Paragraph(
            "Use of proceeds: to fund a portion of the purchase price of the membership interests of Northstar HVAC Services, LLC.",
            s["body"],
        ),
        Paragraph("Financial covenants:", s["h"]),
        Paragraph(
            f"Minimum debt service coverage ratio of {N.COVENANT_DSCR:.2f}x, tested annually on the fiscal year-end financial statements.",
            s["body"],
        ),
        Paragraph(f"Maximum funded debt to EBITDA of {N.COVENANT_MAX_LEVERAGE:.2f}x, tested annually.", s["body"]),
        Paragraph(
            "Cash flow available for debt service is defined as EBITDA less unfinanced capital expenditures, cash taxes, and permitted distributions, in each case for the trailing twelve months.",
            s["body"],
        ),
        Paragraph(
            "Security: first-priority lien on all assets of the borrower and a pledge of membership interests. Personal guaranty of the sponsor principals.",
            s["body"],
        ),
        Paragraph(
            "Conditions: satisfactory quality of earnings review, customer contract review, and insurance certificates.",
            s["body"],
        ),
        PageBreak(),
        Paragraph("Fees and expenses", s["h"]),
        Paragraph(
            "Origination fee of 1.00% of the facility amount payable at closing. Borrower to reimburse lender legal expenses. This term sheet is non-binding and for discussion only.",
            s["body"],
        ),
    ]
    return _pdf(story, "Fictional term sheet")


def build_contract_apex() -> bytes:
    s = _styles()
    story = [
        Paragraph("Master Service Agreement", s["title"]),
        Paragraph("Apex Logistics Park (Customer) and Northstar HVAC Services, LLC (Contractor)", s["h"]),
        Paragraph(N.FICTIONAL_NOTICE, s["notice"]),
        Paragraph(
            "1. Scope. Contractor will provide preventive maintenance, filter replacement, and priority emergency response for the HVAC systems at fourteen (14) distribution facilities listed in Schedule A.",
            s["body"],
        ),
        Paragraph(
            f"2. Term. The initial term of this Agreement begins January 1, 2024 and ends {N.CONTRACT_APEX_TERM_END}. Thereafter this Agreement renews automatically for successive one-year periods unless either party gives written notice of non-renewal at least ninety (90) days before the end of the then-current term.",
            s["body"],
        ),
        Paragraph(
            "3. Termination for convenience. Customer may terminate this Agreement for convenience upon sixty (60) days' written notice to Contractor, without penalty.",
            s["body"],
        ),
        Paragraph(
            f"4. Fees. Customer shall pay an annual base fee of {_m(N.APEX_REVENUE)} for the services in Schedule A, invoiced in equal monthly installments, plus time-and-materials charges for work outside the scope.",
            s["body"],
        ),
        PageBreak(),
        Paragraph(
            "5. Service levels. Contractor will respond to emergency calls within four hours and complete quarterly preventive maintenance visits per the Schedule B checklist.",
            s["body"],
        ),
        Paragraph(
            "6. Price adjustment. Fees may be increased once per year by no more than the change in the regional consumer price index, capped at 3%.",
            s["body"],
        ),
        Paragraph(
            "7. Assignment. This Agreement may not be assigned by Contractor, including by change of control, without Customer's prior written consent, which shall not be unreasonably withheld.",
            s["body"],
        ),
        PageBreak(),
        Paragraph("Schedule A. Facilities: Apex Distribution Centers 1 through 14 (fictional addresses omitted).", s["body"]),
        Paragraph("Signed by the fictional representatives of each party on December 15, 2023.", s["body"]),
    ]
    return _pdf(story, "Fictional Apex MSA")


def build_contract_harbor() -> bytes:
    s = _styles()
    story = [
        Paragraph("HVAC Services Agreement", s["title"]),
        Paragraph("Harbor Medical Campus (Customer) and Northstar HVAC Services, LLC (Contractor)", s["h"]),
        Paragraph(N.FICTIONAL_NOTICE, s["notice"]),
        Paragraph(
            "1. Scope. Contractor will maintain the HVAC and air-handling systems serving two hospital buildings and one outpatient center, including monthly inspections and 24-hour emergency response.",
            s["body"],
        ),
        Paragraph(
            "2. Term. The term of this Agreement is five (5) years, beginning July 1, 2023 and ending June 30, 2028. This Agreement does not renew automatically.",
            s["body"],
        ),
        Paragraph(
            "3. Termination. Either party may terminate this Agreement only for uncured material breach after thirty (30) days' written notice.",
            s["body"],
        ),
        Paragraph(
            "4. Fees. Customer shall pay an annual fee of $1,120,000, invoiced monthly, plus time-and-materials charges for emergency repairs outside the maintenance scope.",
            s["body"],
        ),
        PageBreak(),
        Paragraph(
            "5. Compliance. Contractor will comply with the Customer's infection-control and access policies while on site.",
            s["body"],
        ),
        Paragraph("6. Assignment. Assignment requires Customer's written consent, not to be unreasonably withheld.", s["body"]),
        Paragraph("Signed by the fictional representatives of each party on June 20, 2023.", s["body"]),
    ]
    return _pdf(story, "Fictional Harbor agreement")


def build_financial_statements() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Income Statement"
    bold = Font(bold=True)
    ws.append([N.COMPANY, "", "", ""])
    ws.append(["Income statement, fiscal years ended December 31 (USD). FICTIONAL demonstration data.", "", "", ""])
    ws.append(["Line item", *N.PERIODS])
    for c in ws[3]:
        c.font = bold
    order = [
        ("Revenue", "revenue"),
        ("Cost of goods sold", "cost_of_goods_sold"),
        ("Gross profit", "gross_profit"),
        *[(N.OPEX_LABELS[k], k) for k in N.OPEX_KEYS],
        ("Total operating expenses", "operating_expenses"),
        ("EBITDA", "ebitda"),
        ("Depreciation", "depreciation"),
        ("Amortization", "amortization"),
        ("Operating income", "operating_income"),
        ("Interest expense", "interest_expense"),
        ("Income before tax", "income_before_tax"),
        ("Income tax expense", "income_tax_expense"),
        ("Net income", "net_income"),
    ]
    for label, key in order:
        ws.append([label, *[int(N.STATEMENTS[p][key]) for p in N.PERIODS]])
    ws.column_dimensions["A"].width = 44
    for col in "BCD":
        ws.column_dimensions[col].width = 14
        for row in range(4, 4 + len(order)):
            ws[f"{col}{row}"].number_format = "#,##0"

    bs = wb.create_sheet("Balance Sheet")
    bs.append(["Balance sheet at December 31 (USD). FICTIONAL demonstration data.", "", "", ""])
    bs.append(["Line item", *N.PERIODS])
    for label, vals in [
        ("Cash", [412_000, 468_000, 521_000]),
        ("Accounts receivable", [1_215_000, 1_380_000, 1_560_000]),
        ("Inventory", [286_000, 305_000, 331_000]),
        ("Property and equipment, net", [1_640_000, 1_720_000, 1_805_000]),
        ("Total assets", [3_553_000, 3_873_000, 4_217_000]),
        ("Accounts payable", [612_000, 668_000, 722_000]),
        ("Accrued expenses", [231_000, 254_000, 279_000]),
        ("Vehicle loans", [1_310_000, 1_215_000, 1_085_000]),
        ("Members' equity", [1_400_000, 1_736_000, 2_131_000]),
    ]:
        bs.append([label, *vals])
    bs.column_dimensions["A"].width = 34

    notes = wb.create_sheet("Notes")
    notes.append(["Note", "Text"])
    notes.append(["1", "Owner compensation includes salary and benefits paid to the sole member, $355,000 in FY2024."])
    notes.append(
        [
            "2",
            "Legal and professional fees in FY2024 include a $65,000 settlement of a warranty dispute (fictional Drake v. Northstar). Prior years contain no comparable item.",
        ]
    )
    notes.append(
        ["3", "Temporary labor represents seasonal staffing agency costs incurred in each of the three years presented."]
    )
    notes.append(
        [
            "4",
            "Marketing and advertising in FY2024 includes approximately $95,000 for a brand refresh and website; ongoing digital advertising continued at prior-year levels.",
        ]
    )
    notes.append(["5", "The company is a limited liability company; income tax expense reflects state entity-level taxes only."])
    notes.column_dimensions["B"].width = 120
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_acquisition_model() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Assumptions"
    ws.append(["Northstar HVAC acquisition model (seller-prepared, FICTIONAL)", ""])
    ws.append(["Assumption", "Value"])
    rows = [
        ("Enterprise value", int(N.ENTERPRISE_VALUE)),
        ("Senior debt", int(N.FUNDED_DEBT)),
        ("Sponsor equity", int(N.EQUITY)),
        ("Interest rate", "8.0%"),
        ("Amortization (years)", N.AMORTIZATION_YEARS),
        ("Revenue growth (FY2025-FY2029)", "18.0%"),
        ("Recurring revenue share", "85%"),
        ("Gross margin FY2025", "36.0%"),
        ("Gross margin FY2026", "38.0%"),
        ("Adjusted EBITDA FY2024", int(N.SELLER_ADJUSTED_EBITDA)),
        ("Exit multiple", "6.0x"),
    ]
    for r in rows:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 36
    adj = wb.create_sheet("Adjustments")
    adj.append(["Seller add-back schedule FY2024", "", ""])
    adj.append(["Adjustment", "Amount", "Rationale"])
    for a in N.ADDBACKS:
        adj.append([a["label"], int(a["amount"]), a["seller_rationale"]])
    adj.append(["Total add-backs", int(N.SELLER_ADDBACKS_TOTAL), ""])
    adj.append(["Reported EBITDA", int(N.REPORTED_EBITDA_FY2024), ""])
    adj.append(["Adjusted EBITDA", int(N.SELLER_ADJUSTED_EBITDA), ""])
    adj.column_dimensions["A"].width = 40
    adj.column_dimensions["C"].width = 80
    proj = wb.create_sheet("Projections")
    proj.append(["Seller projections (USD)", "FY2025", "FY2026", "FY2027", "FY2028", "FY2029"])
    rev = Decimal(N.STATEMENTS["FY2024"]["revenue"])
    revs = []
    for _ in range(5):
        rev = rev * Decimal("1.18")
        revs.append(int(rev))
    proj.append(["Revenue", *revs])
    proj.append(["Adjusted EBITDA", *[int(r * Decimal("0.17")) for r in map(Decimal, revs)]])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_customer_csv() -> bytes:
    out = StringIO()
    w = csv.writer(out)
    w.writerow(
        [
            "customer_id",
            "customer_name",
            "segment",
            "revenue_type",
            "fy2024_revenue",
            "fy2023_revenue",
            "contract_start",
            "contract_end",
            "auto_renew",
        ]
    )
    # FY2023 column: scaled so the column totals FY2023 revenue exactly (Apex absorbs rounding).
    factor = Decimal(N.STATEMENTS["FY2023"]["revenue"]) / Decimal(N.CUSTOMER_TOTAL)
    fy23 = [int((Decimal(c[4]) * factor).quantize(Decimal("1"))) for c in N.CUSTOMERS]
    fy23[0] += int(N.STATEMENTS["FY2023"]["revenue"]) - sum(fy23)
    for c, prior in zip(N.CUSTOMERS, fy23, strict=True):
        w.writerow(
            [c[0], c[1], c[2], c[3], c[4], prior, c[6], c[7], c[8]]
            if False
            else [c[0], c[1], c[2], c[3], c[4], prior, c[5], c[6], c[7]]
        )
    return out.getvalue().encode("utf-8")


def build_ground_truth() -> dict:
    st = N.STATEMENTS
    cagr = f.cagr(st["FY2022"]["revenue"], st["FY2024"]["revenue"], 2).value
    ads = f.annual_debt_service(N.FUNDED_DEBT, N.INTEREST_RATE_PCT, N.AMORTIZATION_YEARS, N.PAYMENTS_PER_YEAR).value
    facts: dict[str, Any] = dict(
        base_revenue=st["FY2024"]["revenue"],
        largest_customer_revenue=N.APEX_REVENUE,
        gross_margin_pct=N.GROSS_MARGIN_FY2024_PCT,
        labor_opex=N.LABOR_OPEX_FY2024,
        other_opex=N.OTHER_OPEX_FY2024,
        depreciation_amortization=N.DA_FY2024,
        amortization_years=N.AMORTIZATION_YEARS,
        payments_per_year=N.PAYMENTS_PER_YEAR,
        covenant_dscr_threshold=N.COVENANT_DSCR,
        exit_multiple=f.ev_to_ebitda(N.ENTERPRISE_VALUE, N.VERIFIED_ADJUSTED_EBITDA).value,
        hold_years=5,
    )
    scenarios = {}
    for key, sc in N.SCENARIOS.items():
        out = project(ScenarioInputs(**facts, **sc["assumptions"]))
        scenarios[key] = {
            "assumptions": {k: str(v) for k, v in sc["assumptions"].items()},
            "expected": {
                "year1_revenue": str(out.years[0].revenue),
                "year1_ebitda": str(out.years[0].ebitda),
                "year1_cfads": str(out.years[0].cfads),
                "annual_debt_service": str(out.annual_debt_service),
                "dscr": str(out.dscr),
                "cash_on_cash_pct": str(out.cash_on_cash_pct),
                "irr_pct": None if out.irr_pct is None else str(out.irr_pct),
                "below_threshold": out.dscr is not None and out.dscr < N.COVENANT_DSCR,
                "warning_codes": sorted({w["code"] for w in out.warnings}),
            },
        }
    claims = [
        {
            "key": "cim_revenue_fy2024",
            "document": "cim",
            "page": 3,
            "claim_type": "revenue",
            "metric_key": "revenue",
            "period": "FY2024",
            "claimed_value": "12950000",
            "unit": "usd",
            "expected_status": "supported",
            "expected_verified_value": str(st["FY2024"]["revenue"]),
        },
        {
            "key": "cim_revenue_growth_18pct",
            "document": "cim",
            "page": 3,
            "claim_type": "revenue_growth",
            "metric_key": "cagr",
            "period": "FY2024",
            "claimed_value": "18",
            "unit": "pct",
            "expected_status": "contradicted",
            "expected_verified_value": str(cagr),
        },
        {
            "key": "cim_reported_ebitda_fy2024",
            "document": "cim",
            "page": 3,
            "claim_type": "adjusted_ebitda",
            "metric_key": "ebitda_reported",
            "period": "FY2024",
            "claimed_value": "1640000",
            "unit": "usd",
            "expected_status": "supported",
            "expected_verified_value": str(N.REPORTED_EBITDA_FY2024),
        },
        {
            "key": "cim_adjusted_ebitda_2_10m",
            "document": "cim",
            "page": 3,
            "claim_type": "adjusted_ebitda",
            "metric_key": "ebitda_adjusted_verified",
            "period": "FY2024",
            "claimed_value": "2100000",
            "unit": "usd",
            "expected_status": "contradicted",
            "expected_verified_value": str(N.VERIFIED_ADJUSTED_EBITDA),
        },
        {
            "key": "cim_gross_margin_35pct",
            "document": "cim",
            "page": 4,
            "claim_type": "gross_margin",
            "metric_key": "gross_margin",
            "period": "FY2024",
            "claimed_value": "35.0",
            "unit": "pct",
            "expected_status": "supported",
            "expected_verified_value": str(N.GROSS_MARGIN_FY2024_PCT.quantize(Decimal("0.00000001"))),
        },
        {
            "key": "cim_no_customer_over_10pct",
            "document": "cim",
            "page": 5,
            "claim_type": "customer_concentration",
            "metric_key": "customer_concentration_top1",
            "period": "FY2024",
            "claimed_value": "10",
            "unit": "pct",
            "comparator": "lte",
            "expected_status": "contradicted",
            "expected_verified_value": str(N.LARGEST_CUSTOMER_PCT.quantize(Decimal("0.00000001"))),
        },
        {
            "key": "cim_recurring_85pct",
            "document": "cim",
            "page": 5,
            "claim_type": "recurring_revenue",
            "metric_key": "recurring_revenue_pct",
            "period": "FY2024",
            "claimed_value": "85",
            "unit": "pct",
            "expected_status": "contradicted",
            "expected_verified_value": str(N.RECURRING_PCT.quantize(Decimal("0.00000001"))),
        },
        {
            "key": "cim_churn_below_5pct",
            "document": "cim",
            "page": 5,
            "claim_type": "churn",
            "metric_key": None,
            "period": None,
            "claimed_value": "5",
            "unit": "pct",
            "comparator": "lte",
            "expected_status": "unsupported",
        },
        {
            "key": "cim_apex_term_2027_auto_renew",
            "document": "cim",
            "page": 6,
            "claim_type": "contract_term",
            "metric_key": None,
            "period": None,
            "claimed_value": None,
            "unit": "text",
            "expected_status": "review_required",
        },
        {
            "key": "cim_temp_labor_one_time",
            "document": "cim",
            "page": 7,
            "claim_type": "one_time_expense",
            "metric_key": "opex_temporary_labor_recurring",
            "period": "FY2024",
            "claimed_value": "130000",
            "unit": "usd",
            "expected_status": "contradicted",
        },
        {
            "key": "cim_integration_savings_65k",
            "document": "cim",
            "page": 7,
            "claim_type": "addback",
            "metric_key": None,
            "period": None,
            "claimed_value": "65000",
            "unit": "usd",
            "expected_status": "unsupported",
        },
        {
            "key": "cim_fy2025_forecast_15_3m",
            "document": "cim",
            "page": 8,
            "claim_type": "forecast",
            "metric_key": None,
            "period": "FY2025",
            "claimed_value": "15300000",
            "unit": "usd",
            "expected_status": "unsupported",
        },
        {
            "key": "cim_margin_to_38pct_fy2026",
            "document": "cim",
            "page": 8,
            "claim_type": "margin_improvement",
            "metric_key": None,
            "period": "FY2026",
            "claimed_value": "38",
            "unit": "pct",
            "expected_status": "unsupported",
        },
        {
            "key": "term_sheet_min_dscr_1_25x",
            "document": "debt_term_sheet",
            "page": 1,
            "claim_type": "debt_term",
            "metric_key": "covenant_dscr_threshold",
            "period": None,
            "claimed_value": "1.25",
            "unit": "multiple",
            "expected_status": "supported",
        },
        {
            "key": "term_sheet_rate_8pct",
            "document": "debt_term_sheet",
            "page": 1,
            "claim_type": "debt_term",
            "metric_key": "interest_rate_pct",
            "period": None,
            "claimed_value": "8.0",
            "unit": "pct",
            "expected_status": "supported",
        },
        {
            "key": "term_sheet_amortization_10y",
            "document": "debt_term_sheet",
            "page": 1,
            "claim_type": "debt_term",
            "metric_key": "amortization_years",
            "period": None,
            "claimed_value": "10",
            "unit": "years",
            "expected_status": "supported",
        },
        {
            "key": "model_growth_18pct",
            "document": "acquisition_model",
            "sheet": "Assumptions",
            "claim_type": "forecast",
            "metric_key": "cagr",
            "comparator": "forward",
            "period": "FY2025",
            "claimed_value": "18.0",
            "unit": "pct",
            "expected_status": "unsupported",
        },
        {
            "key": "model_recurring_85pct",
            "document": "acquisition_model",
            "sheet": "Assumptions",
            "claim_type": "recurring_revenue",
            "metric_key": "recurring_revenue_pct",
            "period": "FY2024",
            "claimed_value": "85",
            "unit": "pct",
            "expected_status": "contradicted",
        },
        {
            "key": "apex_contract_term_2026",
            "document": "contract_apex",
            "page": 1,
            "claim_type": "contract_term",
            "metric_key": None,
            "period": None,
            "claimed_value": None,
            "unit": "text",
            "expected_status": "supported",
        },
    ]
    return {
        "schema_version": "1.0",
        "generated_at": FIXED_DATE.isoformat(),
        "fictional_notice": N.FICTIONAL_NOTICE,
        "company": N.COMPANY,
        "files": N.FILES,
        "deal": {
            "purchase_price": str(N.ENTERPRISE_VALUE),
            "purchase_price_basis": "enterprise_value",
            "purchase_date": N.PURCHASE_DATE,
            "debt_amount": str(N.FUNDED_DEBT),
            "equity_amount": str(N.EQUITY),
            "interest_rate_pct": str(N.INTEREST_RATE_PCT),
            "amortization_years": N.AMORTIZATION_YEARS,
            "payments_per_year": N.PAYMENTS_PER_YEAR,
            "covenant_dscr_threshold": str(N.COVENANT_DSCR),
        },
        "income_statement": {p: {k: str(v) for k, v in st[p].items()} for p in N.PERIODS},
        "metrics": {
            "revenue_growth_fy2023_pct": str(f.revenue_growth(st["FY2023"]["revenue"], st["FY2022"]["revenue"]).value),
            "revenue_growth_fy2024_pct": str(f.revenue_growth(st["FY2024"]["revenue"], st["FY2023"]["revenue"]).value),
            "cagr_fy2022_fy2024_pct": str(cagr),
            "cim_claimed_growth_pct": str(N.CIM_REVENUE_GROWTH_CLAIM_PCT),
            "gross_margin_fy2024_pct": str(N.GROSS_MARGIN_FY2024_PCT.quantize(Decimal("0.00000001"))),
            "reported_ebitda_fy2024": str(N.REPORTED_EBITDA_FY2024),
            "seller_addbacks_total": str(N.SELLER_ADDBACKS_TOTAL),
            "seller_adjusted_ebitda": str(N.SELLER_ADJUSTED_EBITDA),
            "accepted_addbacks_total": str(N.ACCEPTED_ADDBACKS_TOTAL),
            "verified_adjusted_ebitda": str(N.VERIFIED_ADJUSTED_EBITDA),
            "largest_customer": "Apex Logistics Park",
            "largest_customer_revenue": str(N.APEX_REVENUE),
            "largest_customer_pct": str(N.LARGEST_CUSTOMER_PCT.quantize(Decimal("0.00000001"))),
            "cim_claimed_max_customer_pct": str(N.CIM_MAX_CUSTOMER_PCT),
            "recurring_revenue": str(N.RECURRING_TOTAL),
            "recurring_revenue_pct": str(N.RECURRING_PCT.quantize(Decimal("0.00000001"))),
            "cim_claimed_recurring_pct": str(N.CIM_RECURRING_PCT),
            "enterprise_value": str(N.ENTERPRISE_VALUE),
            "ev_to_verified_ebitda": str(f.ev_to_ebitda(N.ENTERPRISE_VALUE, N.VERIFIED_ADJUSTED_EBITDA).value),
            "ev_to_seller_ebitda": str(f.ev_to_ebitda(N.ENTERPRISE_VALUE, N.SELLER_ADJUSTED_EBITDA).value),
            "debt_to_verified_ebitda": str(f.debt_to_ebitda(N.FUNDED_DEBT, N.VERIFIED_ADJUSTED_EBITDA).value),
            "annual_debt_service": str(ads),
        },
        "addbacks": [
            {
                "key": a["key"],
                "label": a["label"],
                "amount": str(a["amount"]),
                "expected_decision": a["expected_decision"],
                "expected_rule": a["expected_rule"],
            }
            for a in N.ADDBACKS
        ],
        "scenario_facts": {k: (str(v) if isinstance(v, Decimal) else v) for k, v in facts.items()},
        "scenarios": scenarios,
        "expected_claims": claims,
        "intentional_discrepancies": [
            {
                "key": "revenue_growth",
                "cim_says": f"{N.CIM_REVENUE_GROWTH_CLAIM_PCT}% annual growth",
                "evidence_says": f"{cagr:.1f}% CAGR FY2022-FY2024 (financial statements)",
            },
            {
                "key": "customer_concentration",
                "cim_says": "no customer above 10%",
                "evidence_says": f"Apex Logistics Park {N.LARGEST_CUSTOMER_PCT:.1f}% (customer revenue file)",
            },
            {
                "key": "adjusted_ebitda",
                "cim_says": _m(N.SELLER_ADJUSTED_EBITDA),
                "evidence_says": f"{_m(N.VERIFIED_ADJUSTED_EBITDA)} after add-back review",
            },
            {
                "key": "recurring_revenue",
                "cim_says": f"{N.CIM_RECURRING_PCT}% recurring",
                "evidence_says": f"{N.RECURRING_PCT:.1f}% under maintenance agreements (customer revenue file)",
            },
            {
                "key": "apex_contract_term",
                "cim_says": f"through {N.CIM_APEX_TERM_END} with automatic renewal",
                "evidence_says": f"initial term ends {N.CONTRACT_APEX_TERM_END}; auto-renews annually; customer may terminate for convenience on 60 days' notice",
            },
            {"key": "temporary_labor", "cim_says": "one-time cost", "evidence_says": "recurs in FY2022, FY2023, and FY2024"},
        ],
        "missing_documents": N.MISSING_DOCUMENTS,
        "injection_strings": N.INJECTION_STRINGS,
        "expected_findings": [
            "revenue_growth_contradiction",
            "customer_concentration_contradiction",
            "adjusted_ebitda_contradiction",
            "recurring_revenue_contradiction",
            "covenant_breach_downside",
            "apex_termination_for_convenience",
            "missing_quality_of_earnings",
            "instruction_text_in_cim",
        ],
    }


def generate(out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    builders = {
        N.FILES["cim"]: build_cim,
        N.FILES["debt_term_sheet"]: build_term_sheet,
        N.FILES["contract_apex"]: build_contract_apex,
        N.FILES["contract_harbor"]: build_contract_harbor,
        N.FILES["financial_statements"]: build_financial_statements,
        N.FILES["acquisition_model"]: build_acquisition_model,
        N.FILES["customer_revenue"]: build_customer_csv,
    }
    for name, fn in builders.items():
        path = out_dir / name
        path.write_bytes(fn())
        written.append(path)
    gt = out_dir / N.FILES["ground_truth"]
    gt.write_text(json.dumps(build_ground_truth(), indent=2) + "\n")
    written.append(gt)
    return written
